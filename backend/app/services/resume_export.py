import io
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

# ATS parsers choke on tables, columns, text boxes, and headers/footers holding real content.
# This generator deliberately avoids all of that: one column, plain paragraphs, standard
# heading text (not styled graphics), and no content in the header/footer.


def _set_base_style(doc: Document):
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)


def _heading(doc: Document, text: str):
    p = doc.add_paragraph()
    run = p.add_run(text.upper())
    run.bold = True
    run.font.size = Pt(11.5)
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    # simple bottom border for a visual divider without using a table
    pPr = p._p.get_or_add_pPr()
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "444444")
    pBdr.append(bottom)
    pPr.append(pBdr)


def build_resume_docx(master_resume: dict, tailored_summary: str = "", tailored_bullets: list = None) -> io.BytesIO:
    """master_resume: dict with name/email/phone/linkedin_url/github_url/portfolio_url/
    summary/skills/education/certs/experience (list of {heading, bullets, project_url}).
    If tailored_summary/tailored_bullets are provided (from a scored JD match), they
    replace the generic summary and are surfaced as a "Highlights for this role" section —
    every fact still traces back to the master resume, nothing is invented here."""
    doc = Document()
    for section in doc.sections:
        section.top_margin = Pt(36)
        section.bottom_margin = Pt(36)
        section.left_margin = Pt(50)
        section.right_margin = Pt(50)
    _set_base_style(doc)

    name = master_resume.get("name") or "Your Name"
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(name)
    run.bold = True
    run.font.size = Pt(18)

    contact_bits = [b for b in [
        master_resume.get("email"), master_resume.get("phone"),
        master_resume.get("linkedin_url"), master_resume.get("github_url"),
        master_resume.get("portfolio_url"),
    ] if b]
    if contact_bits:
        p2 = doc.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p2.add_run(" | ".join(contact_bits)).font.size = Pt(9.5)

    summary_text = tailored_summary or master_resume.get("summary") or ""
    if summary_text:
        _heading(doc, "Summary")
        doc.add_paragraph(summary_text)

    if tailored_bullets:
        _heading(doc, "Highlights for This Role")
        for b in tailored_bullets:
            doc.add_paragraph(b, style="List Bullet")

    skills = master_resume.get("skills") or ""
    if skills:
        _heading(doc, "Skills")
        doc.add_paragraph(skills)

    experience = master_resume.get("experience") or []
    if experience:
        _heading(doc, "Experience & Projects")
        for block in experience:
            heading_text = block.get("heading", "")
            project_url = block.get("project_url", "")
            hp = doc.add_paragraph()
            hp.paragraph_format.space_after = Pt(1)
            hrun = hp.add_run(heading_text)
            hrun.bold = True
            if project_url:
                hp.add_run(f"  —  {project_url}").font.size = Pt(9.5)
            for bullet in block.get("bullets", []) or []:
                doc.add_paragraph(bullet, style="List Bullet")

    education = master_resume.get("education") or ""
    certs = master_resume.get("certs") or ""
    cgpa = master_resume.get("cgpa") or ""
    if education or certs or cgpa:
        _heading(doc, "Education & Certifications")
        if education:
            edu_line = education + (f"  —  CGPA: {cgpa}" if cgpa else "")
            doc.add_paragraph(edu_line)
        if certs:
            doc.add_paragraph(certs)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf
